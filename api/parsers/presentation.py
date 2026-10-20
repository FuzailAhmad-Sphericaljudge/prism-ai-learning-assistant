import io
import re
from typing import List, Optional, Union
from pptx import Presentation
from api.parsers.base import ParsedBlock


class PresentationParser:
    """
    Presentation parser for PowerPoint (.pptx) lecture slide decks.
    Preserves 1-based slide numbers, shape/table text hierarchy, and presenter speaker notes.
    """

    @classmethod
    def parse(cls, file_source: Union[str, bytes, io.BytesIO]) -> List[ParsedBlock]:
        """
        Parses a PPTX file from file path, bytes, or stream.
        Returns a list of ParsedBlock objects, one per non-empty slide.
        """
        if isinstance(file_source, bytes):
            stream = io.BytesIO(file_source)
        elif isinstance(file_source, str):
            stream = file_source
        else:
            stream = file_source

        prs = Presentation(stream)
        total_slides = len(prs.slides)
        blocks: List[ParsedBlock] = []

        for idx, slide in enumerate(prs.slides):
            slide_num = idx + 1
            text_elements: List[str] = []

            # 1. Slide Title & Text Frame shapes
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        p_text = paragraph.text.strip()
                        if p_text:
                            text_elements.append(p_text)

                elif shape.has_table:
                    for row in shape.table.rows:
                        row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_cells:
                            text_elements.append(" | ".join(row_cells))

            # 2. Presenter Speaker Notes
            notes_content: Optional[str] = None
            try:
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    notes_raw = slide.notes_slide.notes_text_frame.text.strip()
                    if notes_raw:
                        notes_content = re.sub(r"\s+", " ", notes_raw)
            except Exception:
                notes_content = None

            full_slide_text = "\n".join(text_elements)
            clean_text = re.sub(r"[ \t]+", " ", full_slide_text).strip()

            # If slide has no visible shapes but has notes, still index it
            combined_text = clean_text
            if notes_content and not combined_text:
                combined_text = f"[Speaker Notes]: {notes_content}"

            if combined_text:
                blocks.append(
                    ParsedBlock(
                        text=combined_text,
                        slide=slide_num,
                        notes=notes_content,
                        metadata={
                            "total_slides": total_slides,
                            "slide_index": idx,
                            "has_notes": bool(notes_content),
                        },
                    )
                )

        return blocks
